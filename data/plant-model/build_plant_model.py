"""Builds the Plant -> Entity -> Group mock data model for the plant-level KPIs
in Plant-KPIs.csv (from P1-R1-KPICatalogue).

Design rule: only BASE MEASURES (raw counts, hours, tonnes, money) are stored at
plant level and rolled up. Every KPI is recomputed at each level from the
rolled-up base measures, never averaged across children. That keeps Entity and
Group numbers exactly reconcilable to the plants underneath them.

Outputs (same folder):
  plant_base.csv, plant_kpi.csv, entity_base.csv, entity_kpi.csv,
  group_base.csv, group_kpi.csv, Plant-Entity-Group-Model.xlsx (live formulas)
Run: python3 build_plant_model.py   (seeded, so output is reproducible)
"""
import csv, os, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter as L

random.seed(20261005)
OUT = os.path.dirname(os.path.abspath(__file__))

GROUP = "Group"
# Group -> Entity -> plants (names match the prototype: Entity A1, Plant 02 ...)
HIER = {"Entity A1": ["Plant 01", "Plant 02", "Plant 03"],
        "Entity A2": ["Plant 04", "Plant 05", "Plant 06"]}
PERIODS = [("2026-04", 30), ("2026-05", 31), ("2026-06", 30),
           ("2026-07", 31), ("2026-08", 31), ("2026-09", 30)]

# Base measures: (field, unit, roll-up rule, description)
BASE = [
    ("calendar_hours", "h", "SUM", "Calendar hours in period"),
    ("planned_op_hours", "h", "SUM", "Calendar hours minus planned shutdown"),
    ("run_hours", "h", "SUM", "Hours actually running"),
    ("unplanned_downtime_hours", "h", "SUM", "Unplanned stoppage hours"),
    ("ideal_output_t", "t", "SUM", "Planned op hours x nameplate rate (theoretical max)"),
    ("installed_capacity_t", "t", "SUM", "Calendar hours x nameplate rate"),
    ("planned_production_t", "t", "SUM", "Production plan for the period"),
    ("gross_output_t", "t", "SUM", "All output incl. off-spec"),
    ("good_output_t", "t", "SUM", "On-spec output = reported production"),
    ("feed_input_t", "t", "SUM", "Raw material fed to process"),
    ("feed_contained_t", "t", "SUM", "Valuable content in feed"),
    ("recovered_contained_t", "t", "SUM", "Valuable content recovered to product"),
    ("prod_loss_downtime_t", "t", "SUM", "Unplanned downtime hours x nameplate rate"),
    ("failures", "count", "SUM", "Breakdown events"),
    ("repair_hours", "h", "SUM", "Hours spent repairing breakdowns"),
    ("pm_scheduled", "count", "SUM", "Preventive-maintenance orders due"),
    ("pm_completed", "count", "SUM", "PM orders completed on time"),
    ("variable_cost_k", "cur k", "SUM", "Variable cost incl. fuel (currency, thousands)"),
    ("fuel_cost_k", "cur k", "SUM", "Fuel cost (subset of variable cost)"),
    ("fixed_cost_k", "cur k", "SUM", "Fixed conversion cost"),
    ("water_m3", "m3", "SUM", "Water withdrawn"),
    ("emissions_tco2e", "tCO2e", "SUM", "Scope 1+2 emissions"),
    ("energy_gj", "GJ", "SUM", "Energy consumed"),
    ("production_at_risk_t", "t", "SUM", "Forecast shortfall vs plan for rest of period"),
    ("critical_plant_alerts", "count", "SUM", "Open critical plant-state alerts at period end"),
    ("hours_worked", "h", "SUM", "Employee + contractor hours worked"),
    ("recordable_injuries", "count", "SUM", "OSHA-style recordable injuries"),
    ("severity_incidents", "count", "SUM", "Severity incidents"),
    ("near_misses", "count", "SUM", "Near misses reported"),
    ("critical_safety_incidents", "count", "SUM", "Critical safety incidents"),
    ("env_excursions", "count", "SUM", "Environmental limit excursions"),
    ("env_violations", "count", "SUM", "Environmental violations (regulator-noted)"),
    ("ehs_ca_open", "count", "SUM", "Open EHS corrective actions at period end"),
    ("ehs_ca_overdue", "count", "SUM", "Of which overdue"),
    ("investigations_due", "count", "SUM", "EHS investigations due"),
    ("investigations_completed", "count", "SUM", "EHS investigations completed"),
    ("critical_material_risks", "count", "SUM", "Critical materials with < 7 days cover"),
    ("dispatches_total", "count", "SUM", "Dispatches scheduled"),
    ("dispatches_delayed", "count", "SUM", "Dispatches late"),
    ("permit_days_to_expiry", "days", "MIN", "Days to nearest licence/permit expiry"),
]
BF = [b[0] for b in BASE]

# KPIs: (id, measure, unit, python fn, excel template using {field} refs, formula text)
def d(a, b): return a / b if b else 0.0
KPI = [
    ("OPS-001", "Production vs plan", "%", lambda r: 100*d(r["good_output_t"], r["planned_production_t"]),
     "IFERROR(100*{good_output_t}/{planned_production_t},0)", "good_output_t / planned_production_t x 100"),
    ("OPS-003", "Capacity utilization", "%", lambda r: 100*d(r["good_output_t"], r["installed_capacity_t"]),
     "IFERROR(100*{good_output_t}/{installed_capacity_t},0)", "good_output_t / installed_capacity_t x 100"),
    ("PLT-001", "Asset utilization", "%", lambda r: 100*d(r["run_hours"], r["calendar_hours"]),
     "IFERROR(100*{run_hours}/{calendar_hours},0)", "run_hours / calendar_hours x 100"),
    ("PLT-002", "OEE", "%", lambda r: 100*d(r["good_output_t"], r["ideal_output_t"]),
     "IFERROR(100*{good_output_t}/{ideal_output_t},0)",
     "Availability x Performance x Quality = good_output_t / ideal_output_t x 100"),
    ("PLT-004", "Recovery percentage", "%", lambda r: 100*d(r["recovered_contained_t"], r["feed_contained_t"]),
     "IFERROR(100*{recovered_contained_t}/{feed_contained_t},0)", "recovered_contained_t / feed_contained_t x 100"),
    ("PLT-005", "Yield", "%", lambda r: 100*d(r["good_output_t"], r["feed_input_t"]),
     "IFERROR(100*{good_output_t}/{feed_input_t},0)", "good_output_t / feed_input_t x 100"),
    ("REL-001", "MTBF", "h", lambda r: d(r["run_hours"], r["failures"]),
     "IFERROR({run_hours}/{failures},0)", "run_hours / failures"),
    ("REL-002", "Mean Time to Repair", "h", lambda r: d(r["repair_hours"], r["failures"]),
     "IFERROR({repair_hours}/{failures},0)", "repair_hours / failures"),
    ("REL-003", "Unplanned downtime", "h", lambda r: r["unplanned_downtime_hours"],
     "{unplanned_downtime_hours}", "unplanned_downtime_hours"),
    ("REL-004", "Production loss from downtime", "t", lambda r: r["prod_loss_downtime_t"],
     "{prod_loss_downtime_t}", "prod_loss_downtime_t"),
    ("REL-005", "Preventive-maintenance compliance", "%", lambda r: 100*d(r["pm_completed"], r["pm_scheduled"]),
     "IFERROR(100*{pm_completed}/{pm_scheduled},0)", "pm_completed / pm_scheduled x 100"),
    ("CST-001", "Cost per tonne", "cur/t",
     lambda r: 1000*d(r["variable_cost_k"]+r["fixed_cost_k"], r["good_output_t"]),
     "IFERROR(1000*({variable_cost_k}+{fixed_cost_k})/{good_output_t},0)",
     "(variable_cost_k + fixed_cost_k) x 1000 / good_output_t"),
    ("CST-002", "Variable cost", "cur k", lambda r: r["variable_cost_k"], "{variable_cost_k}", "variable_cost_k"),
    ("CST-003", "Fuel cost", "cur k", lambda r: r["fuel_cost_k"], "{fuel_cost_k}", "fuel_cost_k"),
    ("SUS-001", "Water usage", "m3", lambda r: r["water_m3"], "{water_m3}", "water_m3"),
    ("SUS-002", "Emissions", "tCO2e", lambda r: r["emissions_tco2e"], "{emissions_tco2e}", "emissions_tco2e"),
    ("SUS-003", "Energy intensity", "GJ/t", lambda r: d(r["energy_gj"], r["good_output_t"]),
     "IFERROR({energy_gj}/{good_output_t},0)", "energy_gj / good_output_t"),
    ("SIG-007", "Production at risk", "t", lambda r: r["production_at_risk_t"],
     "{production_at_risk_t}", "production_at_risk_t"),
    ("SIG-008", "Critical plant-state alerts", "count", lambda r: r["critical_plant_alerts"],
     "{critical_plant_alerts}", "critical_plant_alerts"),
    ("SIG-009", "Critical-material shortage risk", "count", lambda r: r["critical_material_risks"],
     "{critical_material_risks}", "critical_material_risks"),
    ("SIG-012", "Dispatch delays", "%", lambda r: 100*d(r["dispatches_delayed"], r["dispatches_total"]),
     "IFERROR(100*{dispatches_delayed}/{dispatches_total},0)", "dispatches_delayed / dispatches_total x 100"),
    ("SIG-021", "Environmental violations", "count", lambda r: r["env_violations"],
     "{env_violations}", "env_violations"),
    ("EHS-001", "TRIR", "per 200k h", lambda r: 200000*d(r["recordable_injuries"], r["hours_worked"]),
     "IFERROR(200000*{recordable_injuries}/{hours_worked},0)", "recordable_injuries x 200,000 / hours_worked"),
    ("EHS-002", "Severity incidents", "count", lambda r: r["severity_incidents"], "{severity_incidents}", "severity_incidents"),
    ("EHS-003", "Near misses", "count", lambda r: r["near_misses"], "{near_misses}", "near_misses"),
    ("EHS-004", "Critical safety incidents", "count", lambda r: r["critical_safety_incidents"],
     "{critical_safety_incidents}", "critical_safety_incidents"),
    ("EHS-005", "Environmental excursions", "count", lambda r: r["env_excursions"], "{env_excursions}", "env_excursions"),
    ("EHS-006", "Open EHS corrective actions", "count", lambda r: r["ehs_ca_open"], "{ehs_ca_open}", "ehs_ca_open"),
    ("EHS-007", "Overdue EHS actions", "count", lambda r: r["ehs_ca_overdue"], "{ehs_ca_overdue}", "ehs_ca_overdue"),
    ("EHS-009", "EHS investigation completion", "%",
     lambda r: 100*d(r["investigations_completed"], r["investigations_due"]) if r["investigations_due"] else 100.0,
     "IF({investigations_due}=0,100,100*{investigations_completed}/{investigations_due})",
     "investigations_completed / investigations_due x 100 (100 if none due)"),
    ("REG-006", "Licence or permit expiry clock", "days", lambda r: r["permit_days_to_expiry"],
     "{permit_days_to_expiry}", "permit_days_to_expiry (MIN across children)"),
]


def plant_row(rate_tph, profile, days, m):
    """One plant-month of physically consistent base measures."""
    ri = lambda a, b: random.randint(a, b)
    ru = random.uniform
    cal = 24 * days
    planned_shut = ri(0, 1) * ri(24, 72) if m in (1, 4) else ri(0, 12)
    pop = cal - planned_shut
    failures = ri(1, 4) + profile["fail"]
    repair = round(failures * ru(3, 9) * profile["mttr"], 1)
    unpl = round(repair + ru(2, 10), 1)                      # downtime >= repair time
    minor = round(pop * ru(0.01, 0.03), 1)
    run = round(pop - unpl - minor, 1)
    ideal = round(pop * rate_tph)
    gross = round(run * rate_tph * ru(0.88, 0.97) * profile["perf"])
    good = round(gross * ru(0.95, 0.99))
    plan = round(pop * rate_tph * ru(0.86, 0.92))
    yld = ru(0.88, 0.95)
    feed = round(good / yld)
    grade = ru(0.18, 0.24)
    fc = round(feed * grade)
    rec = round(fc * ru(0.84, 0.92) * profile["rec"])
    hours_worked = ri(55, 95) * 1000
    pm_s = ri(30, 60)
    inv_due = ri(0, 3)
    disp = ri(120, 260)
    var_cost = round(good * ru(2.6, 3.4) * profile["cost"], 1)   # cur k
    return {
        "calendar_hours": cal, "planned_op_hours": pop, "run_hours": run,
        "unplanned_downtime_hours": unpl, "ideal_output_t": ideal,
        "installed_capacity_t": round(cal * rate_tph), "planned_production_t": plan,
        "gross_output_t": gross, "good_output_t": good, "feed_input_t": feed,
        "feed_contained_t": fc, "recovered_contained_t": rec,
        "prod_loss_downtime_t": round(unpl * rate_tph), "failures": failures,
        "repair_hours": repair, "pm_scheduled": pm_s,
        "pm_completed": pm_s - ri(0, int(pm_s * profile["pm_gap"])),
        "variable_cost_k": var_cost, "fuel_cost_k": round(var_cost * ru(0.22, 0.32), 1),
        "fixed_cost_k": round(ru(900, 1600), 1),
        "water_m3": round(good * ru(1.8, 2.6)), "emissions_tco2e": round(good * ru(0.55, 0.75)),
        "energy_gj": round(good * ru(2.9, 3.6) * profile["cost"]),
        "production_at_risk_t": max(0, round(plan - good * ru(0.97, 1.03))),
        "critical_plant_alerts": ri(0, 2) + profile["fail"],
        "hours_worked": hours_worked, "recordable_injuries": random.choice([0, 0, 0, 1, 1, 2]),
        "severity_incidents": random.choice([0, 0, 0, 0, 1]), "near_misses": ri(2, 12),
        "critical_safety_incidents": random.choice([0] * 9 + [1]),
        "env_excursions": random.choice([0, 0, 0, 1]), "env_violations": random.choice([0] * 7 + [1]),
        "ehs_ca_open": ri(3, 14), "ehs_ca_overdue": 0,  # set below
        "investigations_due": inv_due, "investigations_completed": inv_due - random.choice([0, 0, 0, 1]) if inv_due else 0,
        "critical_material_risks": random.choice([0, 0, 0, 1, 1, 2]),
        "dispatches_total": disp, "dispatches_delayed": round(disp * ru(0.02, 0.09) * profile["perf_disp"]),
        "permit_days_to_expiry": 0,  # set below
    }


def build_plants():
    rows = []
    for ent, plants in HIER.items():
            for p in plants:
                rate = round(random.uniform(90, 260), 1)    # nameplate t/h
                weak = p == "Plant 02"                      # prototype's under-performing plant
                prof = {"fail": 2 if weak else 0, "mttr": 1.4 if weak else 1.0,
                        "perf": 0.86 if weak else 1.0, "rec": 0.96 if weak else 1.0,
                        "cost": 1.12 if weak else 1.0, "pm_gap": 0.25 if weak else 0.08,
                        "perf_disp": 1.8 if weak else 1.0}
                permit = random.randint(160, 420)
                for m, (per, days) in enumerate(PERIODS):
                    r = plant_row(rate, prof, days, m)
                    r["ehs_ca_overdue"] = random.randint(0, r["ehs_ca_open"] // 3)
                    r["permit_days_to_expiry"] = max(0, permit - 30 * m)
                    rows.append({"group": GROUP, "entity": ent, "plant": p,
                                 "period": per, "nameplate_tph": rate, **r})
    return rows


def rollup(rows, keys):
    out = {}
    for r in rows:
        k = tuple(r[x] for x in keys)
        if k not in out:
            out[k] = {x: r[x] for x in keys}
            out[k].update({f: (r[f]) for f in BF})
            continue
        for f, _, rule, _ in BASE:
            out[k][f] = min(out[k][f], r[f]) if rule == "MIN" else out[k][f] + r[f]
    for v in out.values():
        for f in BF:
            if isinstance(v[f], float): v[f] = round(v[f], 1)
    return list(out.values())


def kpis(rows, keys):
    return [{**{x: r[x] for x in keys}, **{k[0]: round(k[3](r), 2) for k in KPI}} for r in rows]


def calc_trace(plant, ent, grp):
    """Every KPI at every level, with its formula, input values and children."""
    import re
    out = []
    levels = [("Plant", plant, lambda r: r["plant"], None),
              ("Entity", ent, lambda r: r["entity"],
               lambda r: [x["plant"] for x in plant if x["entity"] == r["entity"] and x["period"] == r["period"]]),
              ("Group", grp, lambda r: r["group"],
               lambda r: [x["entity"] for x in ent if x["period"] == r["period"]])]
    for kid, name, unit, fn, xl, text in KPI:
        fields = list(dict.fromkeys(re.findall(r"\{(\w+)\}", xl)))
        for lvl, rows, scope, kids in levels:
            for r in rows:
                rule = {f: next(b[2] for b in BASE if b[0] == f) for f in fields}
                out.append({
                    "kpi_id": kid, "measure": name, "unit": unit, "level": lvl,
                    "scope": scope(r), "period": r["period"], "formula": text,
                    "inputs": "; ".join(f"{f}={r[f]}" for f in fields),
                    "calculation": xl.format(**{f: r[f] for f in fields}).replace("IFERROR(", "(").replace(",0)", ")"),
                    "result": round(fn(r), 2),
                    "inputs_built_from": "plant input (mock)" if kids is None else
                        "; ".join(f"{f} = {rule[f]} of [{', '.join(kids(r))}]" for f in fields),
                })
    return out


def write_csv(name, rows):
    with open(os.path.join(OUT, name), "w", newline="") as f:
        w = csv.DictWriter(f, fieldnames=list(rows[0].keys())); w.writeheader(); w.writerows(rows)


def reconcile(plant, ent, grp):
    for f, _, rule, _ in BASE:
        agg = min if rule == "MIN" else sum
        for g in grp:
            p = agg(r[f] for r in plant if r["period"] == g["period"])
            e = agg(r[f] for r in ent if r["period"] == g["period"])
            assert abs(p - g[f]) < 0.05 and abs(e - g[f]) < 0.05, (f, g["period"], p, e, g[f])


def workbook(plant, ent, grp):
    wb = Workbook()
    hdr = Font(bold=True, color="FFFFFF"); fill = PatternFill("solid", fgColor="24406E")
    kfill = PatternFill("solid", fgColor="E8EEF7")

    def sheet(ws, keys, rows, src=None, src_keys=None):
        cols = keys + BF + [k[0] for k in KPI]
        ws.append(cols)
        for c in ws[1]: c.font, c.fill = hdr, fill
        col = {c: L(i + 1) for i, c in enumerate(cols)}
        n = len(rows)
        for i, r in enumerate(rows, start=2):
            line = [r[k] for k in keys]
            for f, _, rule, _ in BASE:
                if src is None:
                    line.append(r[f])
                else:  # live roll-up formula from the child sheet
                    rng = lambda c: f"'{src}'!${src['col'][c]}$2:${src['col'][c]}${src['n']+1}"
                    crit = ",".join(f"{rng(k)},${col[k]}{i}" for k in src_keys)
                    fn = "MINIFS" if rule == "MIN" else "SUMIFS"
                    line.append(f"={fn}({rng(f)},{crit})")
            for k in KPI:
                line.append("=" + k[4].format(**{f: f"{col[f]}{i}" for f in BF}))
            ws.append(line)
        for c in range(len(keys) + len(BF) + 1, len(cols) + 1):
            ws.cell(1, c).fill = PatternFill("solid", fgColor="3E6BB0")
            for rr in range(2, n + 2):
                ws.cell(rr, c).fill = kfill; ws.cell(rr, c).number_format = "0.00"
        ws.freeze_panes = ws.cell(2, len(keys) + 1)
        for i in range(1, len(cols) + 1): ws.column_dimensions[L(i)].width = 14
        return {"col": col, "n": n}

    ws = wb.active; ws.title = "Plant"
    pk = ["group", "entity", "plant", "period", "nameplate_tph"]
    ps = sheet(ws, pk, plant); ps.update(name="Plant")
    # roll-up helper: sheet() formats the source reference via src['col'] and src name
    class Src(dict):
        def __format__(self, spec): return self["name"]
    es = sheet(wb.create_sheet("Entity"), ["group", "entity", "period"], ent,
               Src(ps), ["entity", "period"]); es.update(name="Entity")
    sheet(wb.create_sheet("Group"), ["group", "period"], grp, Src(es), ["group", "period"])

    d = wb.create_sheet("Dictionary")
    d.append(["Type", "Field / KPI ID", "Name", "Unit", "Roll-up rule / formula", "Description"])
    for c in d[1]: c.font, c.fill = hdr, fill
    for f, u, rule, desc in BASE: d.append(["Base measure", f, f, u, rule, desc])
    for k in KPI: d.append(["KPI", k[0], k[1], k[2], k[5], "Recomputed at every level from rolled-up base measures"])
    for c, w in zip("ABCDEF", (14, 26, 34, 12, 60, 50)): d.column_dimensions[c].width = w
    wb.save(os.path.join(OUT, "Plant-Entity-Group-Model.xlsx"))


if __name__ == "__main__":
    plant = build_plants()
    ent = rollup(plant, ["group", "entity", "period"])
    grp = rollup(ent, ["group", "period"])
    reconcile(plant, ent, grp)
    write_csv("plant_base.csv", plant)
    write_csv("plant_kpi.csv", kpis(plant, ["group", "entity", "plant", "period"]))
    write_csv("entity_base.csv", ent)
    write_csv("entity_kpi.csv", kpis(ent, ["group", "entity", "period"]))
    write_csv("group_base.csv", grp)
    write_csv("group_kpi.csv", kpis(grp, ["group", "period"]))
    write_csv("kpi_calculations.csv", calc_trace(plant, ent, grp))
    workbook(plant, ent, grp)
    print(f"plants={len(plant)} rows, entities={len(ent)}, group={len(grp)}; reconciliation OK")
