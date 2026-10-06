"""Bring the "shown on screen" sheets of the KPI workbooks in line with what an Owner screen displays now.

The workbooks record, next to each calculated KPI, the value a screen shows. When a screen is rebuilt those
columns go stale. This script reads the screen as rendered (tools/owner_screen_dump.js) and updates:
  Owner-KPI-Model.xlsx  · Owner Screens (that screen's rows), Calculated vs Shown, Owner KPIs (shown/screens columns)
  Group-KPI-Model.xlsx  · Calculated vs Shown, Group KPIs (shown/screens columns)
Model values (calculated columns, inputs) are not touched.

Usage: python3 tools/sync_owner_workbook.py O-03 P2-O03-CashLiquidity "Cash & Liquidity Dashboard"
"""
import copy, json, re, subprocess, sys
from pathlib import Path
import openpyxl

ROOT = Path(__file__).resolve().parent.parent
KM = ROOT / "data" / "kpi-model"
code, page, title = sys.argv[1], sys.argv[2], sys.argv[3]
items = json.loads(subprocess.check_output(["node", str(ROOT / "tools/owner_screen_dump.js"), page], cwd=ROOT))
on_screen = {}
for it in items:
    if it["element"] == "Card" and it["kpi"]:
        on_screen.setdefault(it["kpi"], it["shown"])


def num(s):
    m = re.search(r"-?[\d,]+(?:\.\d+)?", str(s).replace("−", "-"))
    return float(m.group(0).replace(",", "")) if m else None


def same(calc, shown):
    a, b = num(calc) if not isinstance(calc, (int, float)) else calc, num(shown)
    if a is None or b is None:
        return "Same" if str(calc).strip() == str(shown).strip() else "Different"
    return "Same" if abs(a - b) <= max(0.05, abs(a) * 0.0005) else "Different"


def screens_with(lst, add):
    s = [x.strip() for x in str(lst or "").split(",") if x.strip() and x.strip() != code]
    if add:
        s.append(code)
    order = lambda x: (not x.startswith("O-"), x)
    return ", ".join(sorted(set(s), key=order)) or "—"


def fix_shown_text(txt, add_value):
    """'384 ₹ m (O-03); 4.1 ₹ m (S-11)' → drop the O-03 part, add the new value for O-03."""
    parts = [p.strip() for p in str(txt or "").split(";") if p.strip()]
    keep = []
    for p in parts:
        m = re.match(r"^(.*)\(([^()]*)\)\s*$", p)
        if not m:
            keep.append(p); continue
        scr = [x.strip() for x in m.group(2).split(",") if x.strip() and x.strip() != code]
        if scr:
            keep.append(m.group(1).strip() + " (" + ", ".join(scr) + ")")
    if add_value is not None:
        keep.append(add_value + " (" + code + ")")
    return "; ".join(keep) if keep else "Not on a screen"


def calc_vs_shown(ws, shown_col, same_col, screens_col, calc_col, wsv):
    hdr = [c.value for c in ws[1]]
    S, D, R, C = (hdr.index(x) + 1 for x in (shown_col, same_col, screens_col, calc_col))
    n = 0
    for r in range(2, ws.max_row + 1):
        kid = ws.cell(r, 1).value
        scr = str(ws.cell(r, R).value or "")
        here = kid in on_screen
        was = code in [x.strip() for x in scr.split(",")]
        if not here and not was:
            continue
        rest = [x.strip() for x in scr.split(",") if x.strip() and x.strip() != code]
        if here:
            ws.cell(r, S).value = on_screen[kid]
            ws.cell(r, D).value = same(wsv.cell(r, C).value, on_screen[kid])   # cached value; the sheet cell is a formula
        elif not rest:
            ws.cell(r, S).value = "Not on a screen"
            ws.cell(r, D).value = "—"
        ws.cell(r, R).value = screens_with(scr, here)
        n += 1
    return n


def kpi_sheet(ws, shown_col, screens_col):
    hdr = [c.value for c in ws[1]]
    S, R = hdr.index(shown_col) + 1, hdr.index(screens_col) + 1
    n = 0
    for r in range(2, ws.max_row + 1):
        kid = ws.cell(r, 1).value
        scr = str(ws.cell(r, R).value or "")
        here = kid in on_screen
        if not here and code not in [x.strip() for x in scr.split(",")]:
            continue
        ws.cell(r, S).value = fix_shown_text(ws.cell(r, S).value, on_screen.get(kid))
        ws.cell(r, R).value = screens_with(scr, here)
        n += 1
    return n


def owner_screens(ws, model_row_of):
    rows = [[c.value for c in row] for row in ws.iter_rows(min_row=2, max_row=ws.max_row)]
    style = [copy.copy(c._style) for c in ws[2]]
    first = next((i for i, r in enumerate(rows) if r[1] == code), None)
    rows = [r for r in rows if r[1] != code]
    new = []
    for it in items:
        mv, src = model_row_of(it)
        new.append([None, code, title, it["element"], it["label"], it["kpi"] or None, "Group", it["shown"], mv, src])
    at = first if first is not None else len(rows)
    rows[at:at] = new
    for i, r in enumerate(rows):
        r[0] = i + 1
    ws.delete_rows(2, ws.max_row)
    for i, r in enumerate(rows):
        for j, v in enumerate(r):
            c = ws.cell(i + 2, j + 1, v)
            c._style = copy.copy(style[j])
    ws.auto_filter.ref = "A1:%s%d" % (openpyxl.utils.get_column_letter(len(style)), len(rows) + 1)
    return len(new)


# Model values for the Owner Screens rows: KPI rows from the Owner KPIs sheet, chart and detail rows from Group Inputs
wb = openpyxl.load_workbook(KM / "Owner-KPI-Model.xlsx")
wbv = openpyxl.load_workbook(KM / "Owner-KPI-Model.xlsx", data_only=True)
ok = wbv["Owner KPIs"]
kpi_row = {ok.cell(r, 1).value: r for r in range(2, ok.max_row + 1)}
gi = list(wbv["Group Inputs"].iter_rows(values_only=True))
gi_col = {c: i for i, c in enumerate(gi[1])}
DETAIL = {"Hedged (Sep)": "fx_hedged_k", "Gross FX exposure (Sep)": "fx_exposure_k"}


def model_row_of(it):
    if it["kpi"] and it["kpi"] in kpi_row:
        r = kpi_row[it["kpi"]]
        if it["element"] == "Chart":
            return ("Apr–Sep: " + ", ".join(str(ok.cell(r, c).value) for c in range(6, 12)), "Owner KPIs sheet, row %d (Group, Apr–Sep 2026)" % r)
        return (ok.cell(r, 11).value, "Owner KPIs sheet, row %d (Group, Sep 2026)" % r)
    for k, col in DETAIL.items():
        if it["label"].endswith(k):
            return (round(gi[-1][gi_col[col]] / 1000, 2), "Group Inputs, %s ÷ 1000 (Sep 2026)" % col)
    if it["element"] == "Chart" and "FX" in it["label"]:
        vals = lambda col: ", ".join(str(round(r[gi_col[col]] / 1000, 1)) for r in gi[3:9])
        return ("Exposure " + vals("fx_exposure_k") + " · hedged " + vals("fx_hedged_k"), "Group Inputs, fx_exposure_k and fx_hedged_k ÷ 1000 (Apr–Sep 2026)")
    return (None, "")


n1 = owner_screens(wb["Owner Screens"], model_row_of)
n2 = calc_vs_shown(wb["Calculated vs Shown"], "Shown on the Owner screens (Group scope)", "Same or different", "Screens", "Calculated value, Sep 2026 (P06)", wbv["Calculated vs Shown"])
n3 = kpi_sheet(wb["Owner KPIs"], "Shown on Owner screens today (value and screens)", "Owner screens")
wb.save(KM / "Owner-KPI-Model.xlsx")
g = openpyxl.load_workbook(KM / "Group-KPI-Model.xlsx")
gv = openpyxl.load_workbook(KM / "Group-KPI-Model.xlsx", data_only=True)
n4 = calc_vs_shown(g["Calculated vs Shown"], "Shown on the mockups (Group)", "Same or different", "Screens", "Calculated value, Sep 2026 (P06)", gv["Calculated vs Shown"])
n5 = kpi_sheet(g["Group KPIs"], "Shown on the mockups today (Group)", "Screens")
g.save(KM / "Group-KPI-Model.xlsx")
print("%s: Owner Screens %d rows · Calculated vs Shown %d (Owner) / %d (Group) · KPI sheets %d / %d" % (code, n1, n2, n4, n3, n5))
