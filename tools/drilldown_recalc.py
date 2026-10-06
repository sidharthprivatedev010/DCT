"""Layer 1 of tools/check_drilldowns.js: recompute the model bottom up from KPI-Model.xlsx.

1. Input roll-up: every Entity Inputs cell that the sheet says is a "Sum"/"Lowest" of the entity's
   3 plants is checked numerically against Plant Inputs; every Group Inputs cell that says
   Sum/Lowest/FIRST/SQRTSUMSQ of A1 and A2 is checked against Entity Inputs.
2. KPI values: every formula on Group/A1/A2/PlantNN KPIs sheets is evaluated fresh (xleval)
   from those inputs, P01..P06.
Prints JSON {"inputs": [failures], "kpi": {id: {scope: [P01..P06]}}} for the Node checker.
Usage: python3 tools/drilldown_recalc.py
"""
import json, math, os, sys
import openpyxl

ROOT = os.path.join(os.path.dirname(__file__), "..")
sys.path.insert(0, os.path.join(ROOT, "data/kpi-model"))
from xleval import Book, num  # noqa: E402

wb = openpyxl.load_workbook(os.path.join(ROOT, "data/kpi-model/KPI-Model.xlsx"))
cells = {}
for ws in wb:
    for row in ws.iter_rows():
        for c in row:
            if c.value is not None: cells[(ws.title, c.column, c.row)] = c.value
B = Book(cells, {})
val = lambda s, c, r: B.value(s, c, r)
ENT = {"A1": ["Plant 01", "Plant 02", "Plant 03"], "A2": ["Plant 04", "Plant 05", "Plant 06"]}


def rule_of(note):
    n = (note or "").lower()
    if n.startswith("sum") or n.startswith("sum of"): return "SUM"
    if n.startswith("lowest") or n.startswith("min "): return "MIN"
    if n.startswith("first"): return "FIRST"
    if n.startswith("sqrtsumsq"): return "SQRTSUMSQ"
    return None


def combine(rule, xs):
    xs = [num(x) for x in xs]
    return {"SUM": lambda: sum(xs), "MIN": lambda: min(xs), "FIRST": lambda: xs[0],
            "SQRTSUMSQ": lambda: math.sqrt(sum(x * x for x in xs))}[rule]()


def close(a, b): return abs(a - b) <= max(1e-6, 1e-6 * max(abs(a), abs(b)))


fails = []; checked = 0
# Plant Inputs rows: (plant, period) -> row
pi = wb["Plant Inputs"]; prow = {}
for r in range(4, pi.max_row + 1):
    if pi.cell(r, 3).value: prow[(pi.cell(r, 3).value, pi.cell(r, 4).value)] = r
pcol = {pi.cell(2, c).value: c for c in range(1, pi.max_column + 1)}

ei = wb["Entity Inputs"]; erow = {}
for r in range(4, ei.max_row + 1):
    if ei.cell(r, 2).value: erow[(ei.cell(r, 2).value.replace("Entity ", ""), ei.cell(r, 3).value)] = r
for c in range(5, ei.max_column + 1):
    f, rule = ei.cell(2, c).value, rule_of(ei.cell(3, c).value)
    if not rule or f not in pcol: continue
    for (e, per), r in erow.items():
        kids = [val("Plant Inputs", pcol[f], prow[(p, per)]) for p in ENT[e]]
        got, want = num(val("Entity Inputs", c, r)), combine(rule, kids); checked += 1
        if not close(got, want):
            fails.append({"level": "Entity", "scope": e, "field": f, "period": per, "rule": rule, "got": got, "want": want})

gi = wb["Group Inputs"]; ecol = {ei.cell(2, c).value: c for c in range(1, ei.max_column + 1)}
for c in range(5, gi.max_column + 1):
    f, rule = gi.cell(2, c).value, rule_of(gi.cell(3, c).value)
    if not rule or f not in ecol: continue
    for r in range(4, gi.max_row + 1):
        per = gi.cell(r, 2).value
        kids = [val("Entity Inputs", ecol[f], erow[(e, per)]) for e in ("A1", "A2")]
        got, want = num(val("Group Inputs", c, r)), combine(rule, kids); checked += 1
        if not close(got, want):
            fails.append({"level": "Group", "scope": "Group", "field": f, "period": per, "rule": rule, "got": got, "want": want})

kpi = {}
for sheet in wb.sheetnames:
    if not sheet.endswith(" KPIs"): continue
    scope = sheet[:-5]
    ws = wb[sheet]
    for r in range(2, ws.max_row + 1):
        kid = ws.cell(r, 1).value
        if not kid or not isinstance(kid, str) or "-" not in kid: continue
        out = []
        for c in range(6, 12):
            try: v = val(sheet, c, r); out.append(v if isinstance(v, str) else num(v))
            except Exception as e: out.append("ERR " + str(e))
        kpi.setdefault(kid, {})[scope] = out

json.dump({"inputsChecked": checked, "inputs": fails, "kpi": kpi}, sys.stdout)
